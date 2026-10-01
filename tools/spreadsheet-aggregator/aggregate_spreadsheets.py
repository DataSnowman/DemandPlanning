#!/usr/bin/env python3
"""Append multiple retailer spreadsheets from a (OneDrive) folder into one summary workbook.

Reads every .xlsx in an input folder (skipping any existing summary*.xlsx), validates
that the columns line up, drops blank/total rows, and writes summary<dateRun>.xlsx with a
formatted header and a `Source File` column so each row's origin is preserved.

Usage:
    python aggregate_spreadsheets.py --input "<folder>" [--output "<folder>"]
                                     [--date-format %Y-%m-%d] [--sheet-name "Summary"]
                                     [--strict]

Examples:
    # OneDrive synced folder on Windows
    python aggregate_spreadsheets.py --input "C:\\Users\\me\\OneDrive\\SurfaceCatalogs"

    # Explicit output folder and timestamped file name
    python aggregate_spreadsheets.py --input ".\\spreadsheets" --output ".\\spreadsheets" \
        --date-format %Y-%m-%d_%H%M
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import os
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Row values (case-insensitive) that indicate a total/subtotal row to be dropped.
TOTAL_MARKERS = {
    "total", "totals", "grand total", "subtotal", "sub-total",
    "total units in stock", "average price", "average", "sum",
}


def is_total_row(row: pd.Series) -> bool:
    """Return True if a row looks like a leftover total/subtotal row."""
    non_null = row.dropna()
    if non_null.empty:
        return True  # fully blank row
    first = str(row.iloc[0]).strip().lower()
    if first in TOTAL_MARKERS:
        return True
    # A row with a label in the first cell but mostly empty elsewhere is suspect.
    if first and len(non_null) <= 2 and any(m in first for m in ("total", "average", "sum")):
        return True
    return False


def find_input_files(folder: str) -> list[str]:
    files = sorted(
        f for f in glob.glob(os.path.join(folder, "*.xlsx"))
        if not os.path.basename(f).lower().startswith("summary")
        and not os.path.basename(f).startswith("~$")  # Excel lock files
    )
    return files


def load_one(path: str) -> pd.DataFrame:
    df = pd.read_excel(path)
    df.columns = [str(c).strip() for c in df.columns]
    df = df[~df.apply(is_total_row, axis=1)].copy()
    df.insert(0, "Source File", os.path.basename(path))
    return df


def align_columns(frames: list[pd.DataFrame], strict: bool) -> tuple[list[pd.DataFrame], list[str]]:
    reference = [c for c in frames[0].columns if c != "Source File"]
    warnings: list[str] = []
    for i, df in enumerate(frames[1:], start=2):
        cols = [c for c in df.columns if c != "Source File"]
        if cols != reference:
            missing = [c for c in reference if c not in cols]
            extra = [c for c in cols if c not in reference]
            msg = f"File #{i} columns differ. Missing: {missing or 'none'}; Extra: {extra or 'none'}."
            if strict:
                raise SystemExit(f"ERROR (strict mode): {msg}")
            warnings.append(msg)
    return frames, warnings


def style_output(path: str, sheet_name: str) -> None:
    wb = load_workbook(path)
    ws = wb[sheet_name]
    header_fill = PatternFill("solid", start_color="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    ncols = ws.max_column
    nrows = ws.max_row
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header = str(c.value or "")
        width = max(12, min(48, len(header) + 4))
        ws.column_dimensions[get_column_letter(col)].width = width
    for r in range(2, nrows + 1):
        for col in range(1, ncols + 1):
            ws.cell(row=r, column=col).font = Font(name="Arial", size=10)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Append spreadsheets into summary<dateRun>.xlsx")
    ap.add_argument("--input", required=True, help="Folder containing the source .xlsx files")
    ap.add_argument("--output", default=None, help="Output folder (default: same as --input)")
    ap.add_argument("--date-format", default="%Y-%m-%d",
                    help="strftime format for the <dateRun> suffix (default: %%Y-%%m-%%d)")
    ap.add_argument("--sheet-name", default="Summary", help="Worksheet name in the output file")
    ap.add_argument("--strict", action="store_true",
                    help="Fail if any file's columns do not match the first file")
    args = ap.parse_args()

    in_folder = os.path.abspath(args.input)
    out_folder = os.path.abspath(args.output) if args.output else in_folder
    if not os.path.isdir(in_folder):
        raise SystemExit(f"ERROR: input folder not found: {in_folder}")
    os.makedirs(out_folder, exist_ok=True)

    files = find_input_files(in_folder)
    if len(files) < 2:
        raise SystemExit(f"ERROR: need at least 2 source spreadsheets, found {len(files)} in {in_folder}")

    print(f"Found {len(files)} source file(s):")
    for f in files:
        print(f"  - {os.path.basename(f)}")

    frames = [load_one(f) for f in files]
    frames, warnings = align_columns(frames, args.strict)
    for w in warnings:
        print(f"WARNING: {w}")

    combined = pd.concat(frames, ignore_index=True, sort=False)

    date_run = dt.datetime.now().strftime(args.date_format)
    out_name = f"summary{date_run}.xlsx"
    out_path = os.path.join(out_folder, out_name)

    combined.to_excel(out_path, index=False, sheet_name=args.sheet_name)
    style_output(out_path, args.sheet_name)

    print(f"\nWrote {out_path}")
    print(f"  Rows: {len(combined)}  Columns: {combined.shape[1]}")
    print(f"  Sources combined: {combined['Source File'].nunique()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
